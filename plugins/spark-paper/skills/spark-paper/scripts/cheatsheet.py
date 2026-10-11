#!/usr/bin/env python3
"""cheatsheet.py: build and update the spark-paper reading workbook (.xlsx).

Part of the `spark-paper` skill. One workbook holds the cheat sheets for many
papers: a row per paper on `Papers`, and per-paper rows on `Findings`,
`Methods`, `Figures`, `Terms`, `Discussion`, and `CrossCheck`. The skill
produces one structured RECORD (JSON) per annotated paper; this script upserts
that record into the workbook.

Design rules (each one answers a real failure mode):

* Upsert, never append blindly. Re-annotating a paper replaces that paper's
  rows (matched on the stable paper `id`) and leaves every other paper alone.
* The workbook is the source of truth. Existing rows are read back from the
  visible sheets, so hand edits to other papers survive a rewrite.
* Reader-owned columns ("Notes", "Answer", "Status" on Discussion) are
  carried over onto re-generated rows whenever the row's key still matches.
* The status gate is enforced here, not only in prose: a record that claims
  `status: read` without `cross_check.done` is written as `summary`, with a
  warning. A note earns `read` only after the draft-blind cross-check.
* Every Findings row carries a "Who produced it" label (what the authors
  measured vs. what they cite vs. what they interpret vs. background). Rows
  with a missing or unknown label are rejected, not guessed.
* Text from papers is always stored as text. A cell that begins with `=`, `+`,
  `-` or `@` is never allowed to become a formula.
* Never write over a workbook that is open in an editor (lock-file check), and
  keep one `.prev` copy of the file as it was before each rewrite.

Requires openpyxl (`pip install openpyxl`). Everything else is standard library.

Usage:
	cheatsheet.py new     --workbook PATH
	cheatsheet.py check   RECORD.json
	cheatsheet.py upsert  RECORD.json [RECORD2.json ...] --workbook PATH [--force]
	cheatsheet.py info    --workbook PATH
	cheatsheet.py import-legacy OLD.xlsx --workbook NEW.xlsx

The record format is documented in templates.md (section D).
"""
import argparse
import datetime
import hashlib
import json
import os
import re
import shutil
import sys
import tempfile
import math

try:
	from openpyxl import Workbook, load_workbook
	from openpyxl.formatting.rule import CellIsRule
	from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
	from openpyxl.utils import get_column_letter
	from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
	sys.exit("cheatsheet.py needs openpyxl: pip install openpyxl")

DEFAULT_WORKBOOK = "./papers/Paper_Reading_Notes.xlsx"

# Controlled vocabularies. These are the labels the whole skill uses; the
# drop-downs in the workbook and the validation below both read from here.
STATUSES = ["to-read", "summary", "read"]
PRODUCERS = [
	"MEASURED BY THE AUTHORS",
	"REPORTED FROM CITED WORK",
	"AUTHORS' INTERPRETATION",
	"BACKGROUND STATEMENT",
]
LABELS = ["stated", "inferred", "unspecified"]
KINDS = [
	"Question",
	"Limitation (stated by the authors)",
	"Limitation (evident to a careful reader)",
]
QALMRI = ["Q", "A", "L", "M", "R", "I"]
CHECKS = ["matched", "corrected", "added", "open"]
CLASSES = [
	"value mismatch",
	"draft-only",
	"source-only",
	"emphasis mismatch",
	"label mismatch",
]
DISCUSSION_STATUS = ["open", "resolved", "parked"]
PAPER_TYPES = [
	"research article", "review", "editorial", "method", "benchmark",
	"theory", "preprint", "protocol", "other",
]

TEAL = "1A7F7A"
BODY_FONT = "Calibri"
BODY_SIZE = 10
MAX_ROW_HEIGHT = 409.0  # Excel's hard ceiling for a row
MAX_ROWS = 5000  # validation / formula range depth

_CTRL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


class Col:
	"""One workbook column: header text, width, and how to treat its cells."""

	def __init__(self, header, width, kind="text", user=False, hidden=False):
		self.header = header
		self.width = width
		self.kind = kind  # text | int | link
		self.user = user  # user-owned: preserved across re-annotation
		self.hidden = hidden


# Sheet layouts. Header names are stable identifiers (rows are read back by
# header, so re-ordering columns in the editor does not break an upsert).
SHEETS = [
	("Papers", [
		Col("#", 5, "int"),
		Col("ID", 14),
		Col("Paper", 44),
		Col("DOI", 26, "link"),
		Col("PMID", 11),
		Col("Type", 14),
		Col("Topic", 22),
		Col("Status", 10),
		Col("Cross-check", 18),
		Col("What it sets out to do", 38),
		Col("Summary", 70),
		Col("Main takeaways", 90),
		Col("Zotero", 13, "link"),
		Col("PDF", 9, "link"),
		Col("Tags", 22),
		Col("Updated", 11),
		Col("Notes", 30, user=True),
	]),
	("Findings", [
		Col("#", 5, "int"),
		Col("Paper", 20),
		Col("Where", 26, "link"),
		Col("Finding", 100),
		Col("Number", 30),
		Col("Who produced it", 26),
		Col("Evidence", 14),
		Col("Check", 11),
		Col("Notes", 30, user=True),
		Col("Key", 14, hidden=True),
	]),
	("Methods", [
		Col("#", 5, "int"),
		Col("Paper", 20),
		Col("Item", 26),
		Col("Detail", 90),
		Col("Label", 12),
		Col("Where", 22, "link"),
		Col("Check", 11),
		Col("Notes", 30, user=True),
		Col("Key", 14, hidden=True),
	]),
	("Figures", [
		Col("#", 5, "int"),
		Col("Paper", 20),
		Col("Figure or table", 16),
		Col("What it shows", 70),
		Col("Values read", 50),
		Col("Supports", 16),
		Col("Where", 22, "link"),
		Col("Notes", 30, user=True),
		Col("Key", 14, hidden=True),
	]),
	("Terms", [
		Col("Term", 30),
		Col("Plain definition", 90),
		Col("Paper", 20),
		Col("Where", 22, "link"),
		Col("Analogy (checked)", 40),
		Col("Breaks at", 40),
		Col("Notes", 30, user=True),
		Col("Key", 14, hidden=True),
	]),
	("Discussion", [
		Col("#", 5, "int"),
		Col("Paper", 20),
		Col("Kind", 26),
		Col("Point or question", 100),
		Col("Where", 22, "link"),
		Col("QALMRI", 9),
		Col("Status", 10, user=True),
		Col("Answer", 40, user=True),
		Col("Key", 14, hidden=True),
	]),
	("CrossCheck", [
		Col("#", 5, "int"),
		Col("Paper", 20),
		Col("Field", 20),
		Col("Draft said", 44),
		Col("Verifier said", 44),
		Col("Class", 18),
		Col("Resolution", 44),
		Col("Locus", 22),
		Col("Key", 14, hidden=True),
	]),
]
SHEET_COLS = dict(SHEETS)

# Older headers still read back correctly (reader columns are never dropped).
ALIASES = {"My notes": "Notes", "My answer": "Answer"}


def col_letter(sheet, header):
	"""Column letter of `header` on `sheet`, so formulas never drift from the layout."""
	for ci, col in enumerate(SHEET_COLS[sheet], 1):
		if col.header == header:
			return get_column_letter(ci)
	raise KeyError("%s has no column %r" % (sheet, header))


# ---------------------------------------------------------------- text helpers

def clean(value):
	"""Return `value` as a plain string safe to store in a cell."""
	if value is None:
		return ""
	if isinstance(value, (list, tuple)):
		value = "; ".join(clean(v) for v in value)
	return _CTRL.sub("", str(value)).strip()


def short_id(text):
	"""Make a stable, filesystem-safe paper id from free text."""
	ident = re.sub(r"[^A-Za-z0-9]+", "", clean(text))
	return ident[:32] or "paper"


def row_key(paper_id, *parts):
	"""Stable key tying a row to its paper and its core text."""
	digest = hashlib.sha1("\x1f".join(clean(p) for p in parts).encode("utf-8"))
	return "%s|%s" % (paper_id, digest.hexdigest()[:8])


def key_paper(key):
	"""Recover the paper id from a row key."""
	return clean(key).split("|", 1)[0]


def doi_url(doi):
	"""Resolver URL for a bare DOI, or None when it is not a real DOI."""
	doi = clean(doi)
	for prefix in ("https://doi.org/", "http://doi.org/", "doi:"):
		if doi.lower().startswith(prefix):
			doi = doi[len(prefix):]
	if doi.startswith("10."):
		return "https://doi.org/" + doi
	return None


def bare_doi(doi):
	"""Strip resolver prefixes so the DOI column holds the bare identifier."""
	doi = clean(doi)
	for prefix in ("https://doi.org/", "http://doi.org/", "doi:"):
		if doi.lower().startswith(prefix):
			return doi[len(prefix):]
	return doi


def zotero_select_uri(item_key):
	return "zotero://select/library/items/%s" % item_key if item_key else None


def zotero_pdf_uri(pdf_key, page=None):
	if not pdf_key:
		return None
	uri = "zotero://open-pdf/library/items/%s" % pdf_key
	if page:
		uri += "?page=%d" % int(page)
	return uri


# ------------------------------------------------------------- record -> rows

def validate_record(rec):
	"""Return (errors, warnings) for one paper record. Never mutates it."""
	errors, warnings = [], []
	if not isinstance(rec, dict):
		return ["record is not a JSON object"], []
	if not clean(rec.get("citation")) and not clean(rec.get("short")):
		errors.append("record needs `citation` or `short`")
	status = clean(rec.get("status")) or "summary"
	if status not in STATUSES:
		errors.append("status %r not in %s" % (status, STATUSES))
	cc = rec.get("cross_check") or {}
	if status == "read" and not cc.get("done"):
		warnings.append(
			"status 'read' without cross_check.done: will be written as 'summary'")
	if status in ("summary", "read") and not (rec.get("findings") or []):
		warnings.append("status %r but no findings" % status)
	if cc.get("done") and not clean(cc.get("date")):
		warnings.append("cross_check.done without a date")
	for i, f in enumerate(rec.get("findings") or [], 1):
		if not clean(f.get("where")) and not f.get("page"):
			warnings.append("findings[%d] has no `where` or `page` (every claim needs a locus)" % i)
		if clean(f.get("producer")) not in PRODUCERS:
			errors.append("findings[%d] producer %r not in %s"
				% (i, f.get("producer"), PRODUCERS))
		if not clean(f.get("finding")):
			errors.append("findings[%d] has no `finding` text" % i)
		if clean(f.get("check")) and clean(f.get("check")) not in CHECKS:
			errors.append("findings[%d] check %r not in %s" % (i, f.get("check"), CHECKS))
	for i, m in enumerate(rec.get("methods") or [], 1):
		if clean(m.get("label")) not in LABELS:
			errors.append("methods[%d] label %r not in %s" % (i, m.get("label"), LABELS))
	for i, d in enumerate(rec.get("discussion") or [], 1):
		if clean(d.get("kind")) not in KINDS:
			errors.append("discussion[%d] kind %r not in %s" % (i, d.get("kind"), KINDS))
		if clean(d.get("qalmri")) and clean(d.get("qalmri")) not in QALMRI:
			errors.append("discussion[%d] qalmri %r not in %s" % (i, d.get("qalmri"), QALMRI))
	for i, c in enumerate(rec.get("crosscheck_log") or [], 1):
		if clean(c.get("class")) not in CLASSES:
			errors.append("crosscheck_log[%d] class %r not in %s" % (i, c.get("class"), CLASSES))
	return errors, warnings


def _where(entry, pdf_key):
	"""(display text, url or None) for a row's `where` / `page` fields."""
	text = clean(entry.get("where"))
	page = entry.get("page")
	if not text and page:
		text = "p. %s" % page
	return (text, zotero_pdf_uri(pdf_key, page) if page else None)


def effective_status(rec):
	"""Apply the status gate: `read` requires a completed cross-check."""
	status = clean(rec.get("status")) or "summary"
	if status == "read" and not (rec.get("cross_check") or {}).get("done"):
		return "summary"
	return status


def rows_from_record(rec, paper_id, short):
	"""Expand one record into {sheet: [row dict]} (without paper numbers)."""
	z = rec.get("zotero") or {}
	pdf_key = clean(z.get("pdf_key"))
	out = {name: [] for name, _ in SHEETS}

	for f in rec.get("findings") or []:
		where = _where(f, pdf_key)
		out["Findings"].append({
			"Paper": short, "Where": where,
			"Finding": clean(f.get("finding")), "Number": clean(f.get("number")),
			"Who produced it": clean(f.get("producer")),
			"Evidence": clean(f.get("figure")), "Check": clean(f.get("check")),
			"Key": row_key(paper_id, f.get("finding"), where[0]),
		})
	for m in rec.get("methods") or []:
		where = _where(m, pdf_key)
		out["Methods"].append({
			"Paper": short, "Item": clean(m.get("item")), "Detail": clean(m.get("detail")),
			"Label": clean(m.get("label")), "Where": where, "Check": clean(m.get("check")),
			"Key": row_key(paper_id, m.get("item"), m.get("detail")),
		})
	for g in rec.get("figures") or []:
		where = _where(g, pdf_key)
		out["Figures"].append({
			"Paper": short, "Figure or table": clean(g.get("ref")),
			"What it shows": clean(g.get("shows")), "Values read": clean(g.get("values")),
			"Supports": clean(g.get("supports")), "Where": where,
			"Key": row_key(paper_id, g.get("ref"), g.get("shows")),
		})
	for t in rec.get("terms") or []:
		where = _where(t, pdf_key)
		out["Terms"].append({
			"Term": clean(t.get("term")), "Plain definition": clean(t.get("definition")),
			"Paper": short, "Where": where,
			"Analogy (checked)": clean(t.get("analogy")), "Breaks at": clean(t.get("breaks")),
			"Key": row_key(paper_id, t.get("term")),
		})
	for d in rec.get("discussion") or []:
		where = _where(d, pdf_key)
		out["Discussion"].append({
			"Paper": short, "Kind": clean(d.get("kind")),
			"Point or question": clean(d.get("point")), "Where": where,
			"QALMRI": clean(d.get("qalmri")), "Status": "open",
			"Key": row_key(paper_id, d.get("kind"), d.get("point")),
		})
	for c in rec.get("crosscheck_log") or []:
		out["CrossCheck"].append({
			"Paper": short, "Field": clean(c.get("field")),
			"Draft said": clean(c.get("draft")), "Verifier said": clean(c.get("verifier")),
			"Class": clean(c.get("class")), "Resolution": clean(c.get("resolution")),
			"Locus": clean(c.get("locus")),
			"Key": row_key(paper_id, c.get("field"), c.get("draft"), c.get("verifier")),
		})

	cc = rec.get("cross_check") or {}
	if cc.get("done"):
		cc_text = "done %s (%s diffs)" % (clean(cc.get("date")) or "?", cc.get("disagreements", "?"))
	else:
		cc_text = "not run"
	doi = bare_doi(rec.get("doi"))
	out["Papers"].append({
		"ID": paper_id,
		"Paper": clean(rec.get("citation")) or short,
		"DOI": (doi, doi_url(doi)),
		"PMID": clean(rec.get("pmid")),
		"Type": clean(rec.get("type")),
		"Topic": clean(rec.get("topic")),
		"Status": effective_status(rec),
		"Cross-check": cc_text,
		"What it sets out to do": clean(rec.get("sets_out")),
		"Summary": clean(rec.get("summary")),
		"Main takeaways": clean(rec.get("takeaways")),
		"Zotero": ("Open in Zotero" if z.get("item_key") else "",
			zotero_select_uri(clean(z.get("item_key")))),
		"PDF": ("Open PDF" if pdf_key else "", zotero_pdf_uri(pdf_key)),
		"Tags": clean(rec.get("tags")),
		"Updated": datetime.date.today().isoformat(),
	})
	return out


# ------------------------------------------------------------- workbook <-> data

def _lock_files(path):
	"""Lock files that editors leave next to a workbook that is open."""
	folder, name = os.path.split(os.path.abspath(path))
	return [
		os.path.join(folder, ".~lock.%s#" % name),  # LibreOffice
		os.path.join(folder, "~$%s" % name),  # Microsoft Office
	]


def check_not_open(path, force):
	"""Refuse to rewrite a workbook that an editor currently has open."""
	held = [p for p in _lock_files(path) if os.path.exists(p)]
	if held and not force:
		sys.exit(
			"Refusing to write %s: it looks open in an editor (%s).\n"
			"Save and close it first, or pass --force to override."
			% (path, os.path.basename(held[0])))


def read_workbook(path):
	"""Read every known sheet back as {sheet: [row dict]}.

	Link columns come back as (text, url). Unknown columns are ignored.
	"""
	data = {name: [] for name, _ in SHEETS}
	if not os.path.exists(path):
		return data
	wb = load_workbook(path)
	for name, cols in SHEETS:
		if name not in wb.sheetnames:
			continue
		ws = wb[name]
		index = {}
		for c in ws[1]:
			if c.value is not None:
				index[ALIASES.get(clean(c.value), clean(c.value))] = c.column
		for r in range(2, ws.max_row + 1):
			row, any_value = {}, False
			for col in cols:
				ci = index.get(col.header)
				if ci is None:
					continue
				cell = ws.cell(r, ci)
				val = cell.value
				if col.kind == "link":
					url = cell.hyperlink.target if cell.hyperlink else None
					row[col.header] = (clean(val), url)
					any_value = any_value or bool(clean(val))
				elif col.kind == "int":
					row[col.header] = val if isinstance(val, int) else (
						int(val) if clean(val).isdigit() else None)
				else:
					row[col.header] = clean(val)
					any_value = any_value or bool(clean(val))
			if any_value:
				data[name].append(row)
	return data


def _paper_id_of(sheet, row):
	if sheet == "Papers":
		return clean(row.get("ID"))
	return key_paper(row.get("Key", ""))


def merge_record(data, rec):
	"""Upsert one record into `data` (the dict from read_workbook)."""
	errors, warnings = validate_record(rec)
	if errors:
		raise ValueError("; ".join(errors))
	short = clean(rec.get("short")) or clean(rec.get("citation"))[:16]
	paper_id = short_id(rec.get("id") or short)
	fresh = rows_from_record(rec, paper_id, short)

	for name, cols in SHEETS:
		user_cols = [c.header for c in cols if c.user]
		old = {}
		for row in data[name]:
			if _paper_id_of(name, row) == paper_id:
				old[row.get("Key") if name != "Papers" else paper_id] = row
		kept = [r for r in data[name] if _paper_id_of(name, r) != paper_id]
		for row in fresh[name]:
			k = row.get("Key") if name != "Papers" else paper_id
			prev = old.get(k)
			for uc in user_cols:
				if prev is not None and clean(prev.get(uc)):
					row[uc] = prev[uc]
			if name == "Papers" and prev is not None:
				row["#"] = prev.get("#")
			kept.append(row)
		data[name] = kept
	return paper_id, warnings


def number_papers(data):
	"""Give every paper a stable number and stamp it on every sheet."""
	numbers = {}
	taken = {r.get("#") for r in data["Papers"] if isinstance(r.get("#"), int)}
	nxt = (max(taken) if taken else 0) + 1
	for row in data["Papers"]:
		if not isinstance(row.get("#"), int):
			row["#"] = nxt
			nxt += 1
		numbers[clean(row["ID"])] = row["#"]
	for name, cols in SHEETS:
		if name == "Papers" or "#" not in [c.header for c in cols]:
			continue
		for row in data[name]:
			row["#"] = numbers.get(_paper_id_of(name, row))
	return numbers


def order_rows(data):
	"""Sort by paper number; Terms alphabetically. Python's sort is stable, so
	rows keep the order the record gave them inside each paper."""
	data["Papers"].sort(key=lambda r: r.get("#") or 0)
	for name in ("Findings", "Methods", "Figures", "Discussion", "CrossCheck"):
		data[name].sort(key=lambda r: r.get("#") or 0)
	data["Terms"].sort(key=lambda r: (clean(r.get("Term")).lower(), clean(r.get("Paper"))))


# -------------------------------------------------------------------- styling

def _estimate_height(row, cols):
	"""Rough wrapped-text height. openpyxl cannot autofit, so approximate."""
	lines = 1
	for col in cols:
		val = row.get(col.header)
		text = clean(val[0] if isinstance(val, tuple) else val)
		if not text or col.kind == "int":
			continue
		per_line = max(4, int(col.width * 1.05))
		n = 0
		for para in text.split("\n"):
			n += max(1, math.ceil(len(para) / per_line))
		lines = max(lines, n)
	return 13.0 * lines + 5


def _style_sheet(ws, cols, nrows):
	header_font = Font(name=BODY_FONT, size=BODY_SIZE, bold=True, color="FFFFFF")
	header_fill = PatternFill("solid", fgColor=TEAL)
	wrap_top = Alignment(wrap_text=True, vertical="top")
	body_font = Font(name=BODY_FONT, size=BODY_SIZE)
	link_font = Font(name=BODY_FONT, size=BODY_SIZE, color="0563C1", underline="single")
	for ci, col in enumerate(cols, 1):
		cell = ws.cell(1, ci)
		cell.font, cell.fill = header_font, header_fill
		cell.alignment = Alignment(wrap_text=True, vertical="center")
		letter = get_column_letter(ci)
		ws.column_dimensions[letter].width = col.width
		ws.column_dimensions[letter].hidden = col.hidden
		if col.user:
			cell.fill = PatternFill("solid", fgColor="5B6B7A")  # grey-blue: reader-owned
	ws.row_dimensions[1].height = 30
	ws.freeze_panes = "A2"
	# Printing: landscape, one page wide, header row repeated on every page.
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	ws.print_title_rows = "1:1"
	if nrows:
		ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), nrows + 1)
	for r in range(2, nrows + 2):
		for ci, col in enumerate(cols, 1):
			cell = ws.cell(r, ci)
			cell.font = link_font if (col.kind == "link" and cell.hyperlink) else body_font
			cell.alignment = wrap_top


def _write_sheet(wb, name, cols, rows, clipped):
	"""Write one sheet. Rows too tall for Excel's row-height limit are listed in `clipped`."""
	ws = wb.create_sheet(name)
	for ci, col in enumerate(cols, 1):
		ws.cell(1, ci, col.header)
	for r, row in enumerate(rows, 2):
		for ci, col in enumerate(cols, 1):
			val = row.get(col.header)
			cell = ws.cell(r, ci)
			if col.kind == "link":
				text, url = val if isinstance(val, tuple) else (clean(val), None)
				cell.value = text or None
				if text:
					cell.data_type = "s"
				if url and text:
					cell.hyperlink = url
			elif col.kind == "int":
				cell.value = val if isinstance(val, int) else None
			else:
				text = clean(val)
				cell.value = text or None
				if text:
					cell.data_type = "s"  # never let paper text become a formula
		height = _estimate_height(row, cols)
		if height > MAX_ROW_HEIGHT:
			clipped.append("%s row %d" % (name, r))
		ws.row_dimensions[r].height = min(MAX_ROW_HEIGHT, height)
	_style_sheet(ws, cols, len(rows))
	return ws


def _add_validation(ws, cols, header, list_ref, nrows):
	"""Attach a drop-down from the Lists sheet to one column."""
	for ci, col in enumerate(cols, 1):
		if col.header == header:
			letter = get_column_letter(ci)
			dv = DataValidation(type="list", formula1=list_ref, allow_blank=True,
				showErrorMessage=False)
			ws.add_data_validation(dv)
			dv.add("%s2:%s%d" % (letter, letter, MAX_ROWS))
			return letter
	return None


def _tint(ws, letter, value, rgb):
	ws.conditional_formatting.add(
		"%s2:%s%d" % (letter, letter, MAX_ROWS),
		CellIsRule(operator="equal", formula=['"%s"' % value],
			fill=PatternFill("solid", bgColor=rgb, fgColor=rgb)))


def _lists_sheet(wb):
	"""Controlled vocabularies, used by the drop-downs."""
	ws = wb.create_sheet("Lists")
	lists = [
		("Status", STATUSES), ("Who produced it", PRODUCERS), ("Label", LABELS),
		("Kind", KINDS), ("QALMRI", QALMRI), ("Check", CHECKS), ("Class", CLASSES),
		("Discussion status", DISCUSSION_STATUS), ("Type", PAPER_TYPES),
	]
	refs = {}
	for ci, (title, values) in enumerate(lists, 1):
		ws.cell(1, ci, title)
		for r, v in enumerate(values, 2):
			ws.cell(r, ci, v)
		letter = get_column_letter(ci)
		ws.column_dimensions[letter].width = max(14, max(len(v) for v in values) + 2)
		refs[title] = "Lists!$%s$2:$%s$%d" % (letter, letter, len(values) + 1)
	for ci in range(1, len(lists) + 1):
		c = ws.cell(1, ci)
		c.font = Font(name=BODY_FONT, size=BODY_SIZE, bold=True, color="FFFFFF")
		c.fill = PatternFill("solid", fgColor=TEAL)
	ws.sheet_properties.tabColor = "9AA5AE"
	return refs


def _about_sheet(wb, data):
	"""Legend, status ladder, and live counts (formulas recalc on open)."""
	ws = wb.active
	ws.title = "About"
	title = Font(name=BODY_FONT, size=14, bold=True, color=TEAL)
	bold = Font(name=BODY_FONT, size=BODY_SIZE, bold=True)
	body = Font(name=BODY_FONT, size=BODY_SIZE)
	wrap = Alignment(wrap_text=True, vertical="top")
	ws.column_dimensions["A"].width = 30
	ws.column_dimensions["B"].width = 110
	ws.page_setup.orientation = "landscape"
	ws.page_setup.fitToWidth = 1
	ws.page_setup.fitToHeight = 0
	ws.sheet_properties.pageSetUpPr.fitToPage = True
	rows = [
		("Reading workbook", None, "title"),
		("What this is", "One cheat-sheet row set per paper. Papers is the index; Findings, Methods, "
			"Figures, Terms, Discussion and CrossCheck hold the detail. Each row says where in the paper "
			"it came from. Generated and updated by the spark-paper skill; the paper text is "
			"cross-checked against the source, not guaranteed true.", None),
		("Reader columns", "Grey-blue headers (Notes, Answer, Discussion Status) belong to the reader and "
			"are kept when a paper is re-annotated. Edits elsewhere persist too, except for the rows of "
			"the paper being re-annotated, which are regenerated.", None),
		("Hidden Key column", "Findings, Methods, Figures, Terms, Discussion and CrossCheck keep a hidden "
			"Key column that ties a row to its paper. Do not edit or delete it.", None),
		("Links", "DOI opens the publisher record. Open in Zotero and Open PDF use zotero:// links; a "
			"Where link that has a page number opens the PDF at that page in the Zotero reader.", None),
		("", None, None),
		("Status ladder", None, "head"),
		("to-read", "Pre-read brief only. Nothing here is a finding yet.", None),
		("summary", "Metadata verified, short note, no deep read and no blind cross-check yet.", None),
		("read", "Full read, and the draft-blind cross-check ran and every disagreement is reconciled "
			"or logged. A paper cannot be marked read without it.", None),
		("", None, None),
		("Who produced it", None, "head"),
		("MEASURED BY THE AUTHORS", "A result the authors obtained in this paper.", None),
		("REPORTED FROM CITED WORK", "A result from another paper that this one reports; the cited work, "
			"not this paper, is the evidence.", None),
		("AUTHORS' INTERPRETATION", "What the authors conclude or argue from the evidence.", None),
		("BACKGROUND STATEMENT", "Context stated without a measurement or a specific citation.", None),
		("", None, None),
		("Method labels", None, "head"),
		("stated", "The paper or its supplement says it explicitly.", None),
		("inferred", "Standard practice or derivable, NOT the paper's own word.", None),
		("unspecified", "The paper omits it: a reproducibility gap.", None),
		("", None, None),
		("Counts (live)", None, "head"),
	]
	r = 1
	for left, right, style in rows:
		a, b = ws.cell(r, 1, left or None), ws.cell(r, 2, right)
		a.font = title if style == "title" else bold if (style == "head" or right) else body
		b.font = body
		a.alignment = b.alignment = wrap
		if right:
			ws.row_dimensions[r].height = max(15.0, 13.0 * math.ceil(len(right) / 118.0) + 3)
		r += 1
	def rng(sheet, header):
		letter = col_letter(sheet, header)
		return "%s!%s2:%s%d" % (sheet, letter, letter, MAX_ROWS)

	counts = [("Papers", "=COUNTA(%s)" % rng("Papers", "ID"))]
	for s in STATUSES:
		counts.append(("  " + s, '=COUNTIF(%s,"%s")' % (rng("Papers", "Status"), s)))
	counts.append(("Findings", "=COUNTA(%s)" % rng("Findings", "Finding")))
	for p in PRODUCERS:
		counts.append(("  " + p.lower(), '=COUNTIF(%s,"%s")' % (rng("Findings", "Who produced it"), p)))
	counts += [
		("Method items", "=COUNTA(%s)" % rng("Methods", "Detail")),
		("  unspecified (gaps)", '=COUNTIF(%s,"unspecified")' % rng("Methods", "Label")),
		("Open questions", '=COUNTIF(%s,"open")' % rng("Discussion", "Status")),
		("Last updated", datetime.date.today().isoformat()),
	]
	for left, formula in counts:
		ws.cell(r, 1, left).font = body
		c = ws.cell(r, 2, formula)
		c.font = body
		c.alignment = Alignment(horizontal="left")
		r += 1


def write_workbook(path, data, force=False):
	"""Render `data` to a fresh workbook and swap it in atomically."""
	check_not_open(path, force)
	number_papers(data)
	order_rows(data)
	wb = Workbook()
	_about_sheet(wb, data)
	sheets_ws = {}
	clipped = []
	for name, cols in SHEETS:
		sheets_ws[name] = _write_sheet(wb, name, cols, data[name], clipped)
	refs = _lists_sheet(wb)

	# Drop-downs.
	n = {name: len(data[name]) for name, _ in SHEETS}
	_add_validation(sheets_ws["Papers"], SHEET_COLS["Papers"], "Status", refs["Status"], n["Papers"])
	_add_validation(sheets_ws["Papers"], SHEET_COLS["Papers"], "Type", refs["Type"], n["Papers"])
	prod = _add_validation(sheets_ws["Findings"], SHEET_COLS["Findings"], "Who produced it",
		refs["Who produced it"], n["Findings"])
	_add_validation(sheets_ws["Findings"], SHEET_COLS["Findings"], "Check", refs["Check"], n["Findings"])
	lab = _add_validation(sheets_ws["Methods"], SHEET_COLS["Methods"], "Label", refs["Label"], n["Methods"])
	_add_validation(sheets_ws["Methods"], SHEET_COLS["Methods"], "Check", refs["Check"], n["Methods"])
	_add_validation(sheets_ws["Discussion"], SHEET_COLS["Discussion"], "Kind", refs["Kind"], n["Discussion"])
	_add_validation(sheets_ws["Discussion"], SHEET_COLS["Discussion"], "QALMRI", refs["QALMRI"], n["Discussion"])
	dstat = _add_validation(sheets_ws["Discussion"], SHEET_COLS["Discussion"], "Status",
		refs["Discussion status"], n["Discussion"])
	cls = _add_validation(sheets_ws["CrossCheck"], SHEET_COLS["CrossCheck"], "Class", refs["Class"], n["CrossCheck"])

	# Colour cues: how much weight a row can bear, at a glance.
	if prod:
		_tint(sheets_ws["Findings"], prod, "MEASURED BY THE AUTHORS", "D9EFD9")
		_tint(sheets_ws["Findings"], prod, "REPORTED FROM CITED WORK", "FFF0CC")
		_tint(sheets_ws["Findings"], prod, "AUTHORS' INTERPRETATION", "DCE8F7")
		_tint(sheets_ws["Findings"], prod, "BACKGROUND STATEMENT", "ECECEC")
	if lab:
		_tint(sheets_ws["Methods"], lab, "stated", "D9EFD9")
		_tint(sheets_ws["Methods"], lab, "inferred", "FFF0CC")
		_tint(sheets_ws["Methods"], lab, "unspecified", "F8D7D3")
	pstat = col_letter("Papers", "Status")
	_tint(sheets_ws["Papers"], pstat, "read", "D9EFD9")
	_tint(sheets_ws["Papers"], pstat, "summary", "FFF0CC")
	_tint(sheets_ws["Papers"], pstat, "to-read", "ECECEC")
	if dstat:
		_tint(sheets_ws["Discussion"], dstat, "resolved", "D9EFD9")
	if cls:
		_tint(sheets_ws["CrossCheck"], cls, "draft-only", "F8D7D3")

	folder = os.path.dirname(os.path.abspath(path))
	os.makedirs(folder, exist_ok=True)
	if os.path.exists(path):
		base, ext = os.path.splitext(path)
		shutil.copy2(path, base + ".prev" + ext)
	fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=folder)
	os.close(fd)
	try:
		wb.save(tmp)
		# mkstemp creates 0600; give the file the mode the old one had, or the umask default
		if os.path.exists(path):
			os.chmod(tmp, os.stat(path).st_mode & 0o777)
		else:
			umask = os.umask(0)
			os.umask(umask)
			os.chmod(tmp, 0o666 & ~umask)
		os.replace(tmp, path)
		if clipped:
			print("warning: text taller than the row-height limit (shorten it or widen the "
				"column): %s" % ", ".join(clipped[:8]), file=sys.stderr)
	finally:
		if os.path.exists(tmp):
			os.remove(tmp)


# ------------------------------------------------------------ legacy import

def _legacy_sheet(wb, name):
	"""Rows of a legacy sheet as dicts keyed by its header text."""
	if name not in wb.sheetnames:
		return []
	ws = wb[name]
	heads = [clean(c.value) for c in ws[1]]
	rows = []
	for r in range(2, ws.max_row + 1):
		vals = {h: clean(ws.cell(r, i + 1).value) for i, h in enumerate(heads) if h}
		if any(vals.values()):
			rows.append(vals)
	return rows


def records_from_legacy(path):
	"""Turn a workbook in the older 4-sheet layout into records.

	The older layout carries no cross-check, so every paper comes in as
	`summary` with `cross_check.done = false`. Nothing is promoted to `read`.
	"""
	wb = load_workbook(path, data_only=True)
	papers = _legacy_sheet(wb, "Papers")
	findings = _legacy_sheet(wb, "Findings")
	terms = _legacy_sheet(wb, "Terms")
	disc = _legacy_sheet(wb, "Discussion")
	short_by_num = {}
	for row in findings + disc:
		num = clean(row.get("#"))
		if num and clean(row.get("Paper")):
			short_by_num.setdefault(num, clean(row["Paper"]))
	recs, by_short = [], {}
	for p in papers:
		num = clean(p.get("#"))
		short = short_by_num.get(num) or "Paper %s" % num
		rec = {
			"id": short_id(short), "short": short, "citation": clean(p.get("Paper")),
			"doi": clean(p.get("DOI")), "pmid": clean(p.get("PMID")),
			"type": clean(p.get("Type")), "topic": clean(p.get("Topic")),
			"status": "summary", "cross_check": {"done": False},
			"sets_out": clean(p.get("What it sets out to do")),
			"summary": clean(p.get("Summary")),
			"takeaways": clean(p.get("Main takeaways")),
			"findings": [], "terms": [], "discussion": [],
		}
		recs.append(rec)
		by_short[short] = rec
	for f in findings:
		rec = by_short.get(clean(f.get("Paper")))
		if rec is None:
			continue
		producer = clean(f.get("Who produced it")).upper().replace("’", "'")
		if producer not in PRODUCERS:
			producer = "BACKGROUND STATEMENT"
		rec["findings"].append({
			"where": clean(f.get("Where")), "finding": clean(f.get("Finding")),
			"number": clean(f.get("Number")), "producer": producer,
		})
	for t in terms:
		rec = by_short.get(clean(t.get("Paper")))
		if rec is None:
			continue
		rec["terms"].append({
			"term": clean(t.get("Term")), "definition": clean(t.get("Plain definition")),
			"where": clean(t.get("Where")),
		})
	for d in disc:
		rec = by_short.get(clean(d.get("Paper")))
		if rec is None:
			continue
		kind = clean(d.get("Kind"))
		if kind not in KINDS:
			kind = "Question"
		rec["discussion"].append({
			"kind": kind, "point": clean(d.get("Point or question")),
			"where": clean(d.get("Where")),
		})
	return recs


# -------------------------------------------------------------------- commands

def _load_record(path):
	with open(path, encoding="utf-8") as fh:
		data = json.load(fh)
	return data if isinstance(data, list) else [data]


def cmd_new(args):
	check_not_open(args.workbook, args.force)
	if os.path.exists(args.workbook):
		sys.exit("%s already exists; use upsert to add papers." % args.workbook)
	write_workbook(args.workbook, read_workbook(args.workbook), args.force)
	print("created %s" % args.workbook)


def cmd_check(args):
	bad = 0
	for path in args.records:
		for rec in _load_record(path):
			errors, warnings = validate_record(rec)
			label = clean(rec.get("short") or rec.get("id") or path)
			for w in warnings:
				print("warning [%s]: %s" % (label, w))
			for e in errors:
				print("error   [%s]: %s" % (label, e))
			bad += len(errors)
			if not errors:
				print("ok      [%s]" % label)
	sys.exit(1 if bad else 0)


def cmd_upsert(args):
	data = read_workbook(args.workbook)
	for path in args.records:
		for rec in _load_record(path):
			try:
				paper_id, warnings = merge_record(data, rec)
			except ValueError as exc:
				sys.exit("record %s rejected: %s" % (path, exc))
			for w in warnings:
				print("warning [%s]: %s" % (paper_id, w))
			print("upserted %s" % paper_id)
	write_workbook(args.workbook, data, args.force)
	print("wrote %s (%d papers)" % (args.workbook, len(data["Papers"])))


def cmd_info(args):
	data = read_workbook(args.workbook)
	if not data["Papers"]:
		print("no papers in %s" % args.workbook)
		return
	for p in data["Papers"]:
		pid = clean(p["ID"])
		counts = {s: sum(1 for r in data[s] if _paper_id_of(s, r) == pid)
			for s in ("Findings", "Methods", "Figures", "Terms", "Discussion", "CrossCheck")}
		print("%-3s %-22s %-8s %s" % (p.get("#"), pid, clean(p.get("Status")),
			" ".join("%s=%d" % (k[:4].lower(), v) for k, v in counts.items())))


def cmd_import_legacy(args):
	data = read_workbook(args.workbook)
	for rec in records_from_legacy(args.legacy):
		paper_id, _ = merge_record(data, rec)
		print("imported %s (%d findings)" % (paper_id, len(rec["findings"])))
	write_workbook(args.workbook, data, args.force)
	print("wrote %s" % args.workbook)


def main(argv=None):
	ap = argparse.ArgumentParser(description="Build and update the spark-paper reading workbook.")
	sub = ap.add_subparsers(dest="cmd", required=True)

	def add_wb(p, required=False):
		p.add_argument("--workbook", default=DEFAULT_WORKBOOK, required=required,
			help="target .xlsx (default %s)" % DEFAULT_WORKBOOK)
		p.add_argument("--force", action="store_true",
			help="write even if an editor lock file is present")

	p = sub.add_parser("new", help="create an empty workbook")
	add_wb(p)
	p.set_defaults(fn=cmd_new)
	p = sub.add_parser("check", help="validate record JSON without writing")
	p.add_argument("records", nargs="+")
	p.set_defaults(fn=cmd_check)
	p = sub.add_parser("upsert", help="add or replace papers from record JSON")
	p.add_argument("records", nargs="+")
	add_wb(p)
	p.set_defaults(fn=cmd_upsert)
	p = sub.add_parser("info", help="list papers and row counts")
	add_wb(p)
	p.set_defaults(fn=cmd_info)
	p = sub.add_parser("import-legacy", help="import a workbook in the older 4-sheet layout")
	p.add_argument("legacy")
	add_wb(p)
	p.set_defaults(fn=cmd_import_legacy)

	args = ap.parse_args(argv)
	args.fn(args)


if __name__ == "__main__":
	main()
